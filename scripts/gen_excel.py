#!/usr/bin/env python3
"""Generate an industrial inspection Excel report from a JSON payload.

Usage: python3 gen_excel.py <input.json> <output.xlsx>

The input JSON has the shape:
{
  "generatedAt": "ISO string",
  "windowDays": 30,
  "kpis": {...},
  "inspections": [...],
  "defects": [...],
  "models": [...],
  "defectDistribution": [...],
  "linePerformance": [...]
}
"""
import sys
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, PieChart, Reference

# Industrial palette (no blue/indigo)
C_HEADER = "1F2937"      # dark slate
C_ACCENT = "B45309"      # amber-700
C_ACCENT_LIGHT = "FEF3C7" # amber-100
C_PASS = "16A34A"
C_FAIL = "DC2626"
C_REVIEW = "CA8A04"
C_GRID = "E5E7EB"
C_SUB = "6B7280"

thin = Side(style="thin", color=C_GRID)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

def style_header(cell):
    cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=C_HEADER)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = BORDER

def style_cell(cell, bold=False, color=None, fill=None, align="left"):
    cell.font = Font(name="Calibri", size=10, bold=bold, color=color or "111827")
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
    cell.border = BORDER

def status_color(status):
    return {"PASS": C_PASS, "FAIL": C_FAIL, "REVIEW": C_REVIEW, "PENDING": C_SUB}.get(status, C_SUB)

def build(wb, data):
    kpis = data.get("kpis", {})
    inspections = data.get("inspections", [])
    defects = data.get("defects", [])
    models = data.get("models", [])
    defect_dist = data.get("defectDistribution", [])
    line_perf = data.get("linePerformance", [])

    # ---- Sheet 1: Summary ----
    ws = wb.active
    ws.title = "Executive Summary"
    ws.sheet_view.showGridLines = False
    ws["A1"] = "AUTOMATED INDUSTRIAL DEFECT DETECTION"
    ws["A1"].font = Font(name="Calibri", size=18, bold=True, color=C_HEADER)
    ws.merge_cells("A1:F1")
    ws["A2"] = "Quality Inspection Report — Industry 4.0 Smart Factory"
    ws["A2"].font = Font(name="Calibri", size=11, italic=True, color=C_SUB)
    ws.merge_cells("A2:F2")
    ws["A3"] = f"Generated: {data.get('generatedAt','')}   |   Window: {data.get('windowDays',30)} days"
    ws["A3"].font = Font(name="Calibri", size=9, color=C_SUB)
    ws.merge_cells("A3:F3")

    row = 5
    ws.cell(row=row, column=1, value="KEY PERFORMANCE INDICATORS").font = Font(name="Calibri", size=12, bold=True, color=C_ACCENT)
    row += 1
    kpi_items = [
        ("Total Inspections", kpis.get("total", 0)),
        ("Passed", kpis.get("passed", 0)),
        ("Failed", kpis.get("failed", 0)),
        ("Flagged for Review", kpis.get("review", 0)),
        ("Pass Rate (%)", kpis.get("passRate", 0)),
        ("Average Confidence", kpis.get("avgConfidence", 0)),
        ("Average Latency (ms)", kpis.get("avgLatency", 0)),
        ("Total Defects Detected", kpis.get("totalDefects", 0)),
        ("Critical Defects", kpis.get("criticalCount", 0)),
    ]
    for label, val in kpi_items:
        c1 = ws.cell(row=row, column=1, value=label); style_cell(c1, bold=True, fill=C_ACCENT_LIGHT)
        c2 = ws.cell(row=row, column=2, value=val); style_cell(c2, align="right")
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=3)
        c3 = ws.cell(row=row, column=3, value=val); style_cell(c3, align="right", bold=True)
        # Actually keep simple: put label in A (merged A:B), value in C
        ws.unmerge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        c1 = ws.cell(row=row, column=1, value=label); style_cell(c1, bold=True, fill=C_ACCENT_LIGHT)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        row += 1

    # Defect distribution table
    row += 1
    ws.cell(row=row, column=1, value="DEFECT TYPE DISTRIBUTION").font = Font(name="Calibri", size=12, bold=True, color=C_ACCENT)
    row += 1
    headers = ["Defect Type", "Count"]
    for i, h in enumerate(headers):
        c = ws.cell(row=row, column=i+1, value=h); style_header(c)
    row += 1
    for d in defect_dist:
        c1 = ws.cell(row=row, column=1, value=d.get("name")); style_cell(c1)
        c2 = ws.cell(row=row, column=2, value=d.get("value", 0)); style_cell(c2, align="right")
        row += 1

    # Line performance table
    row += 1
    ws.cell(row=row, column=1, value="PRODUCTION LINE PERFORMANCE").font = Font(name="Calibri", size=12, bold=True, color=C_ACCENT)
    row += 1
    headers = ["Line Code", "Line Name", "Total", "Pass", "Fail", "Pass Rate (%)"]
    for i, h in enumerate(headers):
        c = ws.cell(row=row, column=i+1, value=h); style_header(c)
    row += 1
    for lp in line_perf:
        vals = [lp.get("code"), lp.get("line"), lp.get("total",0), lp.get("pass",0), lp.get("fail",0), round(lp.get("passRate",0),1)]
        for i, v in enumerate(vals):
            c = ws.cell(row=row, column=i+1, value=v); style_cell(c, align="left" if i<2 else "right")
        row += 1

    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 14
    ws.column_dimensions["F"].width = 16

    # ---- Sheet 2: Inspections ----
    ws2 = wb.create_sheet("Inspection Log")
    ws2.sheet_view.showGridLines = False
    cols = ["Inspection Code", "Part ID", "Part Name", "Part Number", "Status", "Severity",
            "Confidence", "Defects", "Model", "Line", "Latency (ms)", "Inspector", "Created At"]
    for i, h in enumerate(cols):
        c = ws2.cell(row=1, column=i+1, value=h); style_header(c)
    for r, insp in enumerate(inspections, start=2):
        vals = [
            insp.get("inspectionCode"), insp.get("partId"), insp.get("partName"), insp.get("partNumber"),
            insp.get("status"), insp.get("severity"), round(insp.get("confidence",0),3), insp.get("defectCount",0),
            insp.get("modelUsed"), (insp.get("productionLine") or {}).get("code") or "—",
            insp.get("processingMs",0), (insp.get("inspector") or {}).get("name") or "system",
            insp.get("createdAt",""),
        ]
        for i, v in enumerate(vals):
            c = ws2.cell(row=r, column=i+1, value=v)
            fill = None
            if i == 4:  # status column
                style_cell(c, bold=True, color="FFFFFF", fill=status_color(vals[4]), align="center")
            elif i == 5:
                style_cell(c, bold=True, color="FFFFFF", fill=status_color_sev(vals[5]), align="center")
            else:
                style_cell(c, align="left" if i in (0,1,2,3,8,11) else "right")
    widths = [16,16,28,16,10,12,12,9,18,10,12,18,22]
    for i, w in enumerate(widths):
        ws2.column_dimensions[get_column_letter(i+1)].width = w
    ws2.freeze_panes = "A2"

    # ---- Sheet 3: Defects ----
    if defects:
        ws3 = wb.create_sheet("Defect Details")
        ws3.sheet_view.showGridLines = False
        cols = ["Inspection Code", "Part", "Defect Type", "Severity", "Confidence", "Location", "Description"]
        for i, h in enumerate(cols):
            c = ws3.cell(row=1, column=i+1, value=h); style_header(c)
        for r, d in enumerate(defects, start=2):
            vals = [d.get("inspectionCode"), d.get("partName"), d.get("type"), d.get("severity"),
                    round(d.get("confidence",0),3), d.get("location"), d.get("description")]
            for i, v in enumerate(vals):
                c = ws3.cell(row=r, column=i+1, value=v)
                if i == 3:
                    style_cell(c, bold=True, color="FFFFFF", fill=status_color_sev(vals[3]), align="center")
                else:
                    style_cell(c, align="left" if i in (0,1,5,6) else "right")
        widths = [16,28,20,12,12,26,40]
        for i, w in enumerate(widths):
            ws3.column_dimensions[get_column_letter(i+1)].width = w
        ws3.freeze_panes = "A2"

    # ---- Sheet 4: Model Comparison ----
    ws4 = wb.create_sheet("Model Benchmarks")
    ws4.sheet_view.showGridLines = False
    cols = ["Model", "Version", "mAP@0.5", "mAP@0.5:0.95", "Precision", "Recall", "F1", "Inference (ms)", "Params (M)", "Status"]
    for i, h in enumerate(cols):
        c = ws4.cell(row=1, column=i+1, value=h); style_header(c)
    for r, m in enumerate(models, start=2):
        vals = [m.get("modelName"), m.get("version"), round(m.get("map50",0),3), round(m.get("map5095",0),3),
                round(m.get("precision",0),3), round(m.get("recall",0),3), round(m.get("f1",0),3),
                round(m.get("inferenceMs",0),1), round(m.get("paramsMillion",0),1), m.get("status")]
        for i, v in enumerate(vals):
            c = ws4.cell(row=r, column=i+1, value=v); style_cell(c, align="left" if i in (0,1,9) else "right")
    widths = [22,12,12,14,12,12,10,14,12,16]
    for i, w in enumerate(widths):
        ws4.column_dimensions[get_column_letter(i+1)].width = w
    ws4.freeze_panes = "A2"

    # Embed a bar chart of defect distribution on summary
    if defect_dist:
        # write a small data block for the chart
        start = ws.max_row + 3
        ws.cell(row=start, column=1, value="Defect").font = Font(bold=True)
        ws.cell(row=start, column=2, value="Count").font = Font(bold=True)
        for i, d in enumerate(defect_dist):
            ws.cell(row=start+1+i, column=1, value=d.get("name"))
            ws.cell(row=start+1+i, column=2, value=d.get("value",0))
        chart = BarChart()
        chart.type = "bar"
        chart.title = "Defect Distribution"
        chart.y_axis.title = "Defect Type"
        chart.x_axis.title = "Count"
        data = Reference(ws, min_col=2, min_row=start, max_row=start+len(defect_dist))
        cats = Reference(ws, min_col=1, min_row=start+1, max_row=start+len(defect_dist))
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.height = 9
        chart.width = 18
        ws.add_chart(chart, f"D{start}")

def status_color_sev(sev):
    return {"NONE": C_PASS, "MINOR": "CA8A04", "MAJOR": "EA580C", "CRITICAL": C_FAIL}.get(sev, C_SUB)

def main():
    if len(sys.argv) < 3:
        print("usage: gen_excel.py <input.json> <output.xlsx>", file=sys.stderr)
        sys.exit(1)
    with open(sys.argv[1], "r") as f:
        data = json.load(f)
    wb = Workbook()
    wb.properties.creator = "Z.ai Industrial Defect Detection"
    wb.properties.title = "Industrial Inspection Report"
    build(wb, data)
    wb.save(sys.argv[2])
    print("OK", sys.argv[2])

if __name__ == "__main__":
    main()
